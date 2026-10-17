import csv
import io
import re
import zipfile
from datetime import date, datetime, time, timedelta
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException
from django import forms
from django.contrib.auth.forms import UserCreationForm
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.utils import timezone
from .models import BorrowPeriod, Room, User
from .password_policy_r16 import validate_new_password
from .school import term_settings


class BookingForm(forms.Form):
    title = forms.CharField(
        label="预约用途",
        max_length=100,
        widget=forms.TextInput(attrs={"placeholder": "例如：小组研讨、班级活动"}),
    )
    date = forms.DateField(
        label="使用日期",
        widget=forms.DateInput(format="%Y-%m-%d", attrs={"type": "date"}),
    )
    period = forms.ModelChoiceField(
        label="借用节次", queryset=BorrowPeriod.objects.all(), empty_label="请选择节次"
    )
    participants = forms.IntegerField(label="使用人数", min_value=1, max_value=2000)

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["period"].label_from_instance = lambda p: (
            f"{p}（{p.start:%H:%M}–{p.end:%H:%M}）"
        )

    def clean(self):
        data = super().clean()
        if "start" in self.data or "end" in self.data:
            raise ValidationError("借用只能选择固定节次，不能提交自定义起止时间。")
        if data.get("date") and data.get("period"):
            period = data["period"]
            data["start_at"] = timezone.make_aware(
                datetime.combine(data["date"], period.start)
            )
            data["end_at"] = timezone.make_aware(
                datetime.combine(data["date"], period.end)
            )
        return data


class StudentBookingForm(BookingForm):
    student_name = forms.CharField(label="借用人姓名（公开显示）", max_length=60)
    student_id = forms.CharField(label="学号 / 工号（不公开）", max_length=40)
    student_department = forms.CharField(label="学院 / 班级", max_length=80)
    student_contact = forms.CharField(
        label="联系方式（手机或校内邮箱）", max_length=100
    )
    edit_password = forms.CharField(
        label="本次预约修改密码",
        min_length=6,
        max_length=4096,
        strip=False,
        widget=forms.PasswordInput,
        help_text="新密码至少6位，必须含英文字母和数字；请保存好密码。不是校园账号密码。",
    )
    edit_password_confirm = forms.CharField(
        label="再次输入本次密码", max_length=4096, strip=False, widget=forms.PasswordInput
    )

    def clean_edit_password(self):
        password = self.cleaned_data["edit_password"]
        validate_new_password(password)
        return password

    def clean(self):
        data = super().clean()
        if data.get("edit_password") != data.get("edit_password_confirm"):
            self.add_error("edit_password_confirm", "两次密码不一致。")
        return data


class FurnitureLoanForm(forms.Form):
    loan_password = forms.CharField(
        label="桌椅借用密码", min_length=6, max_length=4096,
        widget=forms.PasswordInput(attrs={"autocomplete": "new-password"}),
        help_text="自行设置密码；归还时只需学号和该密码。请妥善保存。",
    )
    loan_password_confirm = forms.CharField(
        label="确认借用密码", min_length=6, max_length=4096,
        widget=forms.PasswordInput(attrs={"autocomplete": "new-password"}),
    )
    borrower_name = forms.CharField(label="姓名", max_length=60)
    phone = forms.CharField(
        label="电话", max_length=30, widget=forms.TextInput(attrs={"type": "tel"})
    )
    student_id = forms.CharField(label="学号 / 工号", max_length=40)
    purpose = forms.CharField(label="借用用途", max_length=120, required=False)

    def clean_loan_password(self):
        password = self.cleaned_data["loan_password"]
        validate_new_password(password)
        return password

    def clean_loan_password_confirm(self):
        a = self.cleaned_data.get("loan_password")
        b = self.cleaned_data.get("loan_password_confirm")
        if a is not None and a != b:
            raise forms.ValidationError("两次填写的桌椅借用密码不一致。")
        return b

    def clean_phone(self):
        value = self.cleaned_data["phone"].strip()
        if not re.fullmatch(r"[0-9+()\- ]{6,30}", value) or len(
            re.sub(r"\D", "", value)
        ) < 6:
            raise ValidationError("请输入有效的联系电话。")
        return value

    def clean(self):
        data = super().clean()
        names = self.data.getlist("item_name")
        customs = self.data.getlist("item_custom_name")
        quantities = self.data.getlist("item_quantity")
        if len(names) != len(quantities) or not 1 <= len(names) <= 20:
            raise ValidationError("最多添加20种物品，每项都要选择物品并填写数量。")
        if customs and len(customs) != len(names):
            raise ValidationError("物品列表已失效，请刷新后重新选择。")
        if not customs:
            customs = [""] * len(names)  # old clients/R9 test forms
        from .models import FurnitureStock
        active = set(FurnitureStock.objects.filter(is_active=True).values_list("name", flat=True))
        all_known = set(FurnitureStock.objects.values_list("name", flat=True))
        items = []
        seen = set()
        for selected, custom, raw_quantity in zip(names, customs, quantities):
            selected, custom, raw_quantity = selected.strip(), custom.strip(), raw_quantity.strip()
            if not selected and not raw_quantity and not custom:
                continue
            if not raw_quantity.isdecimal():
                raise ValidationError("借用数量须为整数。")
            quantity = int(raw_quantity)
            if not 1 <= quantity <= 500:
                raise ValidationError("每种物品数量须为1至500。")
            is_other = selected == "__OTHER__"
            if is_other:
                if not (1 <= len(custom) <= 60) or custom == "__OTHER__":
                    raise ValidationError("选择“其他”后，必须填写1至60字的实际物品名称。")
                if custom in all_known or any(custom.casefold() == x.casefold() for x in all_known):
                    raise ValidationError("该物品已有管理员库存记录，请从下拉列表选取；如已停用，请联系管理员。")
                name = custom
            else:
                if selected not in active:
                    raise ValidationError("请选择管理员当前设置的物品；没有的物品请选择“其他”。")
                name = selected
            if name.casefold() in seen:
                raise ValidationError("同一种物品请合并为一行。")
            seen.add(name.casefold())
            entry = {"name": name, "quantity": quantity}
            if is_other:
                # Not inventory managed; distinguish even if staff later adds the same name.
                entry["untracked"] = True
            items.append(entry)
        if not items:
            raise ValidationError("请至少填写一种借用物品。")
        if sum(item["quantity"] for item in items) > 2000:
            raise ValidationError("单次登记物品总数不能超过2000。")
        data["items"] = items
        return data


class BookingEditForm(BookingForm):
    room = forms.ModelChoiceField(label="教室", queryset=Room.objects.all())
    student_name = forms.CharField(
        label="借用人姓名（公开）", max_length=60, required=False
    )
    student_id = forms.CharField(label="学号 / 工号", max_length=40, required=False)
    student_department = forms.CharField(
        label="学院 / 班级", max_length=80, required=False
    )
    student_contact = forms.CharField(label="联系方式", max_length=100, required=False)
    edit_password = forms.CharField(
        label="更换本次密码（留空保留）",
        required=False,
        min_length=6,
        max_length=4096,
        strip=False,
        widget=forms.PasswordInput,
    )
    reason = forms.CharField(label="管理员修改原因", required=False, max_length=300)

    def clean_edit_password(self):
        password = self.cleaned_data.get("edit_password", "")
        if password:
            validate_new_password(password)
        return password

    def __init__(self, *args, rooms, admin=False, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["room"].queryset = rooms
        self.fields["reason"].required = admin
        if not admin:
            self.fields.pop("reason")


class SchoolSettingsForm(forms.Form):
    name = forms.CharField(label="学期名称", max_length=80)
    starts_on = forms.DateField(
        label="学期开始日期",
        widget=forms.DateInput(format="%Y-%m-%d", attrs={"type": "date"}),
    )
    ends_on = forms.DateField(
        label="学期结束日期",
        widget=forms.DateInput(format="%Y-%m-%d", attrs={"type": "date"}),
    )
    first_week = forms.IntegerField(
        label="开始日期所在的起始周数", min_value=1, max_value=30, initial=1
    )
    booking_open = forms.BooleanField(
        label="开放本学期借用", required=False, initial=True
    )

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for p in BorrowPeriod.objects.all():
            for suffix, label in [("start", "实际开始"), ("end", "实际结束")]:
                self.fields[f"p{p.number}_{suffix}"] = forms.TimeField(
                    label=f"{p.group} · {p} · {label}",
                    widget=forms.TimeInput(format="%H:%M", attrs={"type": "time"}),
                )


class RoomForm(forms.ModelForm):
    class Meta:
        model = Room
        fields = [
            "name",
            "campus",
            "building",
            "floor",
            "capacity",
            "room_type",
            "equipment",
            "department",
            "is_active",
        ]


class AdministratorCreateForm(UserCreationForm):
    class Meta(UserCreationForm.Meta):
        model = User
        fields = ("username", "display_name", "department")


class AdministratorEditForm(forms.ModelForm):
    password1 = forms.CharField(
        label="重设密码（留空保留）",
        required=False,
        max_length=128,
        widget=forms.PasswordInput,
    )
    password2 = forms.CharField(
        label="再次输入新密码",
        required=False,
        max_length=128,
        widget=forms.PasswordInput,
    )

    class Meta:
        model = User
        fields = ("display_name", "department", "is_active")

    def clean(self):
        data = super().clean()
        password = data.get("password1", "")
        if password != data.get("password2", ""):
            self.add_error("password2", "两次密码不一致。")
        if password:
            validate_password(password, self.instance)
        return data


ALIASES = {
    "教室": "room",
    "教室编号": "room",
    "上课教室": "room",
    "教室名称": "room",
    "地点": "room",
    "课程": "title",
    "课程名称": "title",
    "课程名": "title",
    "教师": "teacher",
    "任课教师": "teacher",
    "授课教师": "teacher",
    "日期": "date",
    "上课日期": "date",
    "开始时间": "start",
    "结束时间": "end",
    "起始周": "week_start",
    "开始周": "week_start",
    "结束周": "week_end",
    "终止周": "week_end",
    "单双周": "week_pattern",
    "周模式": "week_pattern",
    "周次": "weeks",
    "星期": "weekday",
    "星期几": "weekday",
    "周几": "weekday",
    "节次": "period",
    "节数": "period",
    "上课节次": "period",
    "排除日期": "excluded_dates",
    "校区": "campus",
    "楼栋": "building",
    "教学楼": "building",
    "楼层": "floor",
    "人数": "capacity",
    "容量": "capacity",
    "教室类型": "room_type",
    "类型": "room_type",
    "首周周一": "term_start",
}


def parse_weekday(value):
    value = str(value).strip().replace("星期", "").replace("周", "")
    return (
        int(value)
        if value.isdigit()
        else {
            "一": 1,
            "二": 2,
            "三": 3,
            "四": 4,
            "五": 5,
            "六": 6,
            "日": 7,
            "天": 7,
        }.get(value, 0)
    )


def week_spec(value):
    text = str(value).strip().replace("（", "(").replace("）", ")")
    parity = "odd" if "单" in text else "even" if "双" in text else "all"
    text = re.sub(r"\([^)]*\)|[周单双]", "", text)
    weeks = set()
    for part in re.split("[,，;；]", text):
        m = re.fullmatch(r"\s*(\d+)\s*(?:[-~至]\s*(\d+))?\s*", part)
        if not m:
            raise ValueError
        a, b = int(m[1]), int(m[2] or m[1])
        if not 1 <= a <= b <= 30:
            raise ValueError
        weeks.update(range(a, b + 1))
    return weeks, parity


def course_times(row):
    if row.get("period"):
        row = {
            **row,
            "period": {
                "第一节": "1",
                "第二节": "2",
                "第三节": "3",
                "第四节": "4",
                "第五节": "5",
            }.get(row["period"], row["period"]),
        }
        m = re.fullmatch(
            r"(?:第)?\s*([1-5])\s*(?:[-~至]\s*([1-5]))?\s*节?", row["period"]
        )
        if not m:
            raise ValueError("课程节次须为 1–5 或连续范围，例如 1-2；不连续请分行")
        a, b = int(m[1]), int(m[2] or m[1])
        periods = {p.number: p for p in BorrowPeriod.objects.all()}
        if a > b or a not in periods or b not in periods:
            raise ValueError("借用节次尚未设置完整")
        return periods[a].start, periods[b].end
    return time.fromisoformat(row["start"]), time.fromisoformat(row["end"])


def metadata(row):
    result = {k: row[k] for k in ("campus", "building") if row.get(k)}
    for key in result:
        if len(result[key]) > 60:
            raise ValueError("校区或楼栋名称过长")
    for key in ("floor", "capacity"):
        if row.get(key):
            value = int(row[key])
            maximum = 99 if key == "floor" else 2000
            if not (0 if key == "floor" else 1) <= value <= maximum:
                raise ValueError("楼层或容量超出范围")
            result[key] = value
    if row.get("room_type"):
        value = {
            "教室": "lecture",
            "普通教室": "lecture",
            "多媒体教室": "lecture",
            "实验室": "lab",
            "实验教室": "lab",
            "seminar": "lecture",
        }.get(row["room_type"], row["room_type"])
        if value not in ("lecture", "lab"):
            raise ValueError("类型请填写教室或实验室")
        result["room_type"] = value
    return result


class CourseImportForm(forms.Form):
    file = forms.FileField(
        label="正式课程课表（Excel XLSX / CSV）",
        widget=forms.FileInput(attrs={"accept": ".xlsx,.csv"}),
    )
    confirm = forms.BooleanField(label="直接导入（不勾选时仅预览验证）", required=False)

    def clean_file(self):
        upload = self.cleaned_data["file"]
        if upload.size > 1024 * 1024:
            raise ValidationError("文件不能超过 1 MB。")
        term = term_settings()
        try:
            raw = upload.read()
            if upload.name.lower().endswith(".xlsx"):
                from openpyxl import load_workbook

                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    if (
                        len(archive.infolist()) > 2000
                        or sum(x.file_size for x in archive.infolist())
                        > 10 * 1024 * 1024
                    ):
                        raise ValidationError("XLSX 解压内容过大。")
                wb = load_workbook(
                    io.BytesIO(raw), read_only=True, data_only=True, keep_links=False
                )
                try:
                    values = wb.worksheets[0].iter_rows(values_only=True)
                    header = list(next(values, ()))
                    records = []
                    for i, row in enumerate(values):
                        if i >= 500:
                            raise ValidationError("最多导入 500 行。")
                        if not any(x is not None for x in row):
                            continue

                        def cell(v):
                            if isinstance(v, datetime):
                                return (
                                    v.date().isoformat()
                                    if v.time() == time.min
                                    else v.isoformat()
                                )
                            if isinstance(v, (date, time)):
                                return v.isoformat()
                            return "" if v is None else str(v).strip()

                        records.append([cell(v) for v in row])
                finally:
                    wb.close()
            else:
                parsed = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
                header = next(parsed, [])
                records = []
                for i, row in enumerate(parsed):
                    if i >= 500:
                        raise ValidationError("最多导入 500 行。")
                    if not any(v.strip() for v in row):
                        continue
                    records.append(row)
            header = [
                ALIASES.get(str(v).strip(), str(v).strip()) if v is not None else ""
                for v in header
            ]
            while header and not header[-1]:
                header.pop()
            if len(set(header)) != len(header):
                raise ValidationError("表头重复，请勿同时填写同义列。")
            if not {"room", "title"} <= set(header) or not (
                "period" in header or {"start", "end"} <= set(header)
            ):
                raise ValidationError(
                    "缺少教室、课程名称或节次/起止时间列，请下载模板。"
                )
            recurring = "date" not in header
            if recurring and not (
                {"weekday"} <= set(header)
                and ("weeks" in header or {"week_start", "week_end"} <= set(header))
            ):
                raise ValidationError("学期课表还需星期和周次（或起始周、结束周）。")
            result = []
            for number, values in enumerate(records, 2):
                while len(values) > len(header) and not values[-1]:
                    values.pop()
                if len(values) != len(header):
                    raise ValueError(f"第 {number} 行列数与表头不一致")
                row = {k: str(v).strip() for k, v in zip(header, values)}
                if not 1 <= len(row["room"]) <= 40 or not 1 <= len(row["title"]) <= 100:
                    raise ValueError(f"第 {number} 行教室或课程名称为空/过长")
                start, end = course_times(row)
                if not time(0) <= start < end <= time(23, 59):
                    raise ValueError(f"第 {number} 行起止时间无效")
                teacher = row.get("teacher", "")
                if len(teacher) > 60:
                    raise ValueError("任课教师名称过长")
                extra = {"room": row["room"], "room_meta": metadata(row)}
                if not recurring:
                    day = date.fromisoformat(row["date"])
                    if not term.starts_on <= day <= term.ends_on:
                        raise ValueError(f"第 {number} 行不在本学期范围")
                    result.append(
                        {
                            **extra,
                            "title": row["title"],
                            "teacher": teacher,
                            "start_at": timezone.make_aware(
                                datetime.combine(day, start)
                            ),
                            "end_at": timezone.make_aware(datetime.combine(day, end)),
                        }
                    )
                    continue
                weekday = parse_weekday(row["weekday"])
                if not 1 <= weekday <= 7:
                    raise ValueError(f"第 {number} 行星期无效")
                weeks, parity = (
                    week_spec(row["weeks"])
                    if row.get("weeks")
                    else (
                        set(range(int(row["week_start"]), int(row["week_end"]) + 1)),
                        "all",
                    )
                )
                if not weeks or min(weeks) < 1 or max(weeks) > 30:
                    raise ValueError("周次须为 1–30")
                pattern = {
                    "每周": "all",
                    "单周": "odd",
                    "双周": "even",
                    "单": "odd",
                    "双": "even",
                }.get(
                    row.get("week_pattern", ""), row.get("week_pattern", "") or parity
                )
                if pattern not in ("all", "odd", "even") or (
                    parity != "all" and pattern != parity
                ):
                    raise ValueError("单双周标记冲突或无效")
                anchor = (
                    date.fromisoformat(row["term_start"])
                    if row.get("term_start")
                    else term.week_one_monday
                )
                if anchor != term.week_one_monday:
                    raise ValueError("首周日期与管理员的学期设置不一致")
                excluded = [
                    date.fromisoformat(x.strip()).isoformat()
                    for x in re.split("[;；]", row.get("excluded_dates", ""))
                    if x.strip()
                ]
                excluded += [
                    (anchor + timedelta(weeks=w - 1, days=weekday - 1)).isoformat()
                    for w in range(min(weeks), max(weeks) + 1)
                    if w not in weeks
                ]
                result.append(
                    {
                        **extra,
                        "rule": {
                            "title": row["title"],
                            "teacher": teacher,
                            "term_start": anchor,
                            "week_start": min(weeks),
                            "week_end": max(weeks),
                            "week_pattern": pattern,
                            "weekday": weekday,
                            "start": start,
                            "end": end,
                            "excluded_dates": sorted(set(excluded)),
                        },
                    }
                )
            if not result:
                raise ValidationError("没有课程数据。")
            return result
        except (
            UnicodeDecodeError,
            csv.Error,
            zipfile.BadZipFile,
            KeyError,
            OSError,
            ValueError,
            ParseError,
            DefusedXmlException,
            InvalidFileException,
        ) as exc:
            raise ValidationError(f"课表格式错误：{exc}") from exc
